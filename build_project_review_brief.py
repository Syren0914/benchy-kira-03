from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "Benchy_Kira_03_Project_Review_Brief.docx"
IMAGE = ROOT / "docs" / "images" / "model-preview.png"

NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F6F7"
MID_GRAY = "606060"
LIGHT_BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "5")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), LIGHT_BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245B8A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(29)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(8)

for style_name, size, before, after in (("Heading 1", 16, 15, 6), ("Heading 2", 12.5, 10, 4)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# First page
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Benchy Kira 03 Engineering Project Review Brief")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(13)
r = p.add_run("Prepared for Project Review")
r.bold = True
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Prepared by Erdene Batbayar\n10 September 2026")
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor.from_string(MID_GRAY)

if IMAGE.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(IMAGE), width=Inches(3.5))
    p = doc.add_paragraph("Figure 1  Enclosure model preview from the supplied 3MF file")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    for run in p.runs:
        run.font.size = Pt(8.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

doc.add_heading("Purpose of this review", level=1)
doc.add_paragraph(
    "I am asking for feedback on the engineering direction and documentation of Benchy, my Kira 03 bench power supply project. The project repurposes a functional computer power supply that is no longer suitable for a modern desktop computer. My goal is to turn it into a useful bench tool and publish enough information for other makers to understand, reproduce, and improve the design."
)
doc.add_paragraph(
    "The enclosure model, project photographs, preliminary parts list, documentation structure, and a clearly marked preliminary connection schematic are available. Exact as-built wiring, the Arduino Nano program, relay behavior, exact component variants, and measured performance still need to be completed before I describe the project as build-ready. I would value feedback on those gaps and on the tests needed to support the final design."
)

doc.add_heading("Project links", level=1)
p = doc.add_paragraph()
p.add_run("Public repository  ").bold = True
add_hyperlink(p, "github.com/Syren0914/benchy-kira-03", "https://github.com/Syren0914/benchy-kira-03")
p = doc.add_paragraph()
p.add_run("Public photo gallery  ").bold = True
add_hyperlink(p, "Google Drive photo folder", "https://drive.google.com/drive/folders/1a5vY3R4KkWWrwnvo-AKPjcpfRcpAucLN")

p = doc.add_heading("Project motivation", level=1)
p.paragraph_format.page_break_before = True
doc.add_paragraph(
    "Desktop computers have become more power-hungry, so a power supply that was adequate ten years ago may no longer meet the demands of a modern computer. That does not mean the older unit has stopped working or has become useless. Rather than discard functional hardware, I repurposed it as the basis of a bench power supply. Benchy extends the useful life of the original supply and turns it into a practical source of power for electronics work, prototyping, and testing."
)
doc.add_paragraph(
    "This project lets me explore engineering reuse as a design problem. A successful conversion requires more than placing an old power supply in a new case. I need to document the available rails, safe operating limits, control behavior, output connections, cooling, mechanical fit, and test results so that the finished tool is useful and understandable."
)
doc.add_paragraph(
    "I also added pixel-art details to give the enclosure a distinctive, friendly identity. I wanted Benchy to be visually engaging as well as useful, rather than another plain workshop instrument."
)

doc.add_heading("Project overview", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(1.45), Inches(2.15), Inches(3.15)]
headers = ["Area", "Current design", "Information still needed"]
for idx, cell in enumerate(table.rows[0].cells):
    cell.width = widths[idx]
    cell.text = headers[idx]
    set_cell_shading(cell, NAVY)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
set_repeat_table_header(table.rows[0])
rows = [
    ("Source hardware", "Reused computer power supply", "Exact model, rail ratings, connector map, condition, and compatibility checks"),
    ("Adjustable output", "SK120 module", "Manufacturer, revision, verified input and output limits, configuration, and thermal behavior"),
    ("Control", "Arduino Nano and relay", "Sketch, board variant, pin map, relay model, switched circuit, and startup state"),
    ("Enclosure", "PLA 3MF model with M2 and M3 fasteners", "Print orientation and profile, screw quantities and lengths, tolerances, and thermal suitability"),
    ("User interface", "Display, rotary control, fixed-voltage labels, and output connectors", "Connector part numbers, polarity, output function, and operating instructions"),
]
for row_index, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for idx, value in enumerate(values):
        cells[idx].width = widths[idx]
        cells[idx].text = value
        cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[idx])
        if row_index % 2 == 0:
            set_cell_shading(cells[idx], PALE_BLUE)
        for paragraph in cells[idx].paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(9)
set_table_borders(table)

doc.add_heading("Design and learning goals", level=1)
add_bullet(doc, "Recover useful value from functional hardware that no longer meets its original application.")
add_bullet(doc, "Integrate mechanical design, additive manufacturing, power conversion, embedded control, and user-interface decisions in one physical system.")
add_bullet(doc, "Measure the finished system and separate manufacturer ratings from limits demonstrated by the complete assembly.")
add_bullet(doc, "Create public documentation that another person can follow without relying on private explanations from me.")

p = doc.add_heading("Current project assets", level=1)
p.paragraph_format.page_break_before = True
add_bullet(doc, "A public GitHub repository containing the project overview, build-guide structure, preliminary parts list, release checklist, and contribution guidance.")
add_bullet(doc, "The supplied power supply v.2.3mf enclosure file and its embedded model preview.")
add_bullet(doc, "A public gallery of forty exterior photographs, including front, rear, top, and angled views.")
add_bullet(doc, "A preliminary component inventory: PLA, M2 and M3 screws, an SK120 module, a reused computer power supply, an Arduino Nano, and a relay.")
add_bullet(doc, "A preliminary color connection schematic showing the intended low-voltage power paths and identifying unconfirmed Nano and relay wiring as TBC.")

doc.add_heading("Engineering status", level=1)
doc.add_paragraph(
    "The prototype enclosure and exterior are documented, but the project is still an engineering draft. The current public documentation intentionally avoids guessing electrical connections or performance. The following work is needed before another person should build or energize the design from the published files."
)
for item in [
    "Identify the exact computer power supply, SK120, Arduino Nano, relay, connectors, fasteners, and wiring materials.",
    "Confirm the preliminary schematic and complete a conductor-by-conductor wiring schedule that states connector orientation and polarity.",
    "Publish the Arduino source code and define the relay state during startup, reset, normal operation, and failure.",
    "Record validated 3D-print settings, part orientation, assembly order, screw sizes, and fit adjustments.",
    "Define and perform no-load, loaded-output, combined-load, display-accuracy, protection, ripple, and thermal tests.",
    "Ask another person to review or reproduce the design using only the released documentation.",
]:
    add_number(doc, item)

doc.add_heading("Safety and scope", level=1)
doc.add_paragraph(
    "The photographed unit includes a mains-style inlet, so the electrical design and commissioning process require careful review. The public draft does not establish the input protection, protective-earth arrangement, insulation, output isolation, fire suitability of the enclosure, or safe service procedure. These items must be verified for the actual power supply and construction before the project is presented as a repeatable build."
)
doc.add_paragraph(
    "I will keep the enclosure closed during ordinary powered demonstrations and will document internal construction only while the unit is safely disconnected. Electrical ratings in the final report will be tied to the exact components and test conditions rather than inferred from labels or vendor maximum values."
)

doc.add_heading("Feedback requested", level=1)
doc.add_paragraph("I would appreciate your comments on the following questions:")
questions = [
    "Is the project objective clear and appropriately scoped for an engineering project?",
    "What additional requirements should I define before finalizing the design?",
    "Does the proposed documentation separate confirmed facts, assumptions, and unverified limits clearly enough?",
    "Which electrical safety checks and load tests should be mandatory before public release?",
    "What measurements would best demonstrate that the converted supply performs reliably for its intended uses?",
    "What changes would improve the enclosure, cooling, controls, connector layout, or serviceability?",
    "What evidence should I include to demonstrate learning, design iteration, and responsible reuse?",
]
for q in questions:
    add_bullet(doc, q)

p = doc.add_heading("Reviewer feedback", level=1)
p.paragraph_format.page_break_before = True
doc.add_paragraph("Reviewer name ____________________________________    Date ____________________")

feedback = doc.add_table(rows=1, cols=3)
feedback.alignment = WD_TABLE_ALIGNMENT.CENTER
feedback.autofit = False
feedback_widths = [Inches(2.0), Inches(1.15), Inches(3.75)]
for idx, value in enumerate(("Review area", "Assessment", "Comments and recommendations")):
    cell = feedback.rows[0].cells[idx]
    cell.width = feedback_widths[idx]
    cell.text = value
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=130, bottom=130)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
set_repeat_table_header(feedback.rows[0])

for row_index, area in enumerate((
    "Problem definition and motivation",
    "Technical design and integration",
    "Testing and evidence",
    "Safety and responsible practice",
    "Documentation and reproducibility",
    "Overall project direction",
), start=1):
    cells = feedback.add_row().cells
    cells[0].text = area
    cells[1].text = ""
    cells[2].text = "\n\n"
    for idx, cell in enumerate(cells):
        cell.width = feedback_widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=130, bottom=130)
        if row_index % 2 == 0:
            set_cell_shading(cell, PALE_BLUE)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(8.8 if idx == 1 else 9.2)
                if idx == 1:
                    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
set_table_borders(feedback)

doc.add_heading("Priority recommendations", level=2)
for prompt in ("1  ", "2  ", "3  "):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.add_run(prompt + "____________________________________________________________________________")

doc.add_heading("Additional comments", level=2)
for _ in range(5):
    p = doc.add_paragraph("________________________________________________________________________________")
    p.paragraph_format.space_after = Pt(7)

doc.core_properties.title = "Benchy Kira 03 Engineering Project Review Brief"
doc.core_properties.subject = "Review of an open source bench power supply project"
doc.core_properties.author = "Erdene Batbayar"
doc.core_properties.keywords = "Benchy, Kira 03, bench power supply, reuse, engineering project"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
