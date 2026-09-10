from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "Benchy_Kira_03_Assembly_Guide.docx"
MODEL_IMAGE = ROOT / "docs" / "images" / "model-preview.png"
SCHEMATIC_IMAGE = ROOT / "hardware" / "connection-schematic.png"

NAVY = "17365D"
BLUE = "2F6EAD"
PALE_BLUE = "EAF1F8"
PALE_YELLOW = "FFF2CC"
MID_GRAY = "5B6573"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = tc_pr.find(qn("w:shd"))
    if fill is None:
        fill = OxmlElement("w:shd")
        tc_pr.append(fill)
    fill.set(qn("w:fill"), color)


def margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color="B7C7D9", size="5"):
    props = table._tbl.tblPr
    element = props.first_child_found_in("w:tblBorders")
    if element is None:
        element = OxmlElement("w:tblBorders")
        props.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), size)
        line.set(qn("w:color"), color)
        element.append(line)


def add_link(paragraph, text, url):
    part = paragraph.part
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def step(doc, number, title, text, checks=()):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f"{number}. {title}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(3)
    for check in checks:
        item = doc.add_paragraph()
        item.paragraph_format.left_indent = Inches(0.18)
        item.paragraph_format.space_after = Pt(2)
        item.add_run("☐ ").bold = True
        item.add_run(check)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for style_name, size, color in (("Title", 26, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 12.5, BLUE)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

header = section.header.paragraphs[0]
header.text = "BENCHY — KIRA 03  |  ASSEMBLY GUIDE"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Open-source prototype • Documentation revision 0.2 • 10 September 2026")
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Benchy — Kira 03")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = RGBColor.from_string(NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prototype Mechanical Assembly Guide")
r.bold = True
r.font.size = Pt(17)
r.font.color.rgb = RGBColor.from_string(BLUE)

if MODEL_IMAGE.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(MODEL_IMAGE), width=Inches(3.55))
    p = doc.add_paragraph("Enclosure preview from power supply v.2.3mf")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.cell(0, 0)
shade(cell, PALE_YELLOW)
margins(cell, top=130, bottom=130)
p = cell.paragraphs[0]
r = p.add_run("CURRENT RELEASE STATUS")
r.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)
p.add_run(
    "  This guide supports printing, dry-fitting, and mechanical assembly of the documented prototype. "
    "The exact hardware variants, screw lengths, internal wiring, firmware, and validated print profile have not yet been released. "
    "Do not use this document as electrical wiring or commissioning instructions."
)

doc.add_heading("Project purpose", level=1)
doc.add_paragraph(
    "Benchy gives a functional older computer power supply a second life as a workshop bench supply. "
    "The enclosure combines the reused supply with an SK120 module, an Arduino Nano, a relay, a display, controls, and output connectors. "
    "Pixel-art details give the practical tool a friendly, distinctive identity."
)
p = doc.add_paragraph()
p.add_run("Project files and updates: ").bold = True
add_link(p, "github.com/Syren0914/benchy-kira-03", "https://github.com/Syren0914/benchy-kira-03")

doc.add_heading("Before you begin", level=1)
bullet(doc, "Download one matching revision of the 3MF model, documentation, bill of materials, and schematic.")
bullet(doc, "Read the entire guide and inspect the model before printing or purchasing hardware.")
bullet(doc, "Keep the computer power supply disconnected during every printing, fitting, and mechanical step.")
bullet(doc, "Do not open or modify the mains-powered circuitry inside the computer power supply.")
bullet(doc, "Have the electrical design, protective measures, wiring, and commissioning reviewed by a qualified person.")

doc.add_page_break()
doc.add_heading("Confirmed materials and components", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(2.05), Inches(1.0), Inches(3.55)]
for idx, text in enumerate(("Item", "Quantity", "Release note")):
    cell = table.rows[0].cells[idx]
    cell.width = widths[idx]
    cell.text = text
    shade(cell, NAVY)
    margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
items = [
    ("PLA enclosure", "1 set", "Use power supply v.2.3mf; slicing settings remain to be validated."),
    ("Computer power supply", "1", "Use the actual tested unit; model and rail ratings must be recorded."),
    ("SK120 module", "1", "Confirm the exact variant and terminal orientation."),
    ("Arduino Nano", "1", "Exact board, mounting, pin map, and program remain to be documented."),
    ("Relay", "1", "Exact relay module, mounting, and electrical function remain to be documented."),
    ("M2 and M3 screws", "As fitted", "Measure and record lengths, head styles, quantities, and mounting locations."),
    ("Panel hardware", "As fitted", "Display, control, inlet/switch, output terminals, and any mating connectors."),
]
for row_number, values in enumerate(items, start=1):
    cells = table.add_row().cells
    for idx, value in enumerate(values):
        cells[idx].width = widths[idx]
        cells[idx].text = value
        margins(cells[idx])
        cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_number % 2 == 0:
            shade(cells[idx], PALE_BLUE)
        for run in cells[idx].paragraphs[0].runs:
            run.font.size = Pt(8.6)
borders(table)

doc.add_heading("Printing and preparation", level=1)
step(doc, 1, "Inspect the 3MF model", "Open power supply v.2.3mf in your slicer and confirm that every intended enclosure part is present. Check overall scale and compare openings with the exact hardware you will install.", (
    "The model is loaded at the intended scale.",
    "Display, connector, control, switch, ventilation, and fastener features are visible.",
    "No geometry errors or unintended loose bodies are reported by the slicer.",
))
step(doc, 2, "Choose and record a print profile", "The supplied 3MF does not contain a validated printer profile. Select settings appropriate for your printer and PLA, then record nozzle size, layer height, walls, infill, supports, orientation, temperatures, and estimated material use. Print small fit checks first when practical.", (
    "Critical openings were checked against measured hardware.",
    "Support material can be removed without damaging vents or pixel-art details.",
    "The selected profile is saved with the project revision.",
))
step(doc, 3, "Inspect the printed parts", "Allow the parts to cool, remove supports carefully, and clear debris from holes and ventilation slots. Reject cracked, warped, softened, or poorly bonded parts.", (
    "Parts sit flat and mate without force.",
    "Ventilation openings are clear.",
    "Fastener holes are intact and have not been enlarged excessively.",
))

doc.add_heading("Prepare the hardware", level=1)
step(doc, 4, "Identify every component", "Photograph the labels and connector faces of the actual computer PSU, SK120, Nano, relay, display, switch/inlet, controls, and output terminals. Record the manufacturer and model where available.", (
    "No component is identified only by appearance or wire color.",
    "Connector orientation and polarity are recorded separately from this mechanical guide.",
))
step(doc, 5, "Sort and measure fasteners", "Group the M2 and M3 screws by length and head style. Test each screw by hand in its intended mounting feature while the component is unpowered. A screw must hold securely without bottoming out or reaching circuitry.", (
    "Each mounting location has a recorded screw size and length.",
    "Washers, nuts, inserts, or spacers are present where the design requires them.",
))

doc.add_heading("Mechanical assembly", level=1)
step(doc, 6, "Dry-fit the enclosure", "Assemble the empty printed shell loosely. Confirm that mating faces align, the enclosure stands securely, and all panels can be removed again. Do not force warped parts into alignment with screws.", (
    "The frame is stable on a flat surface.",
    "Panel seams close evenly.",
    "Pixel-art surfaces and labels face outward in the intended orientation.",
))
step(doc, 7, "Fit the front-panel hardware", "With the enclosure open and the equipment disconnected, insert the display, rotary control, buttons, fixed-output terminals, adjustable-output terminals, and upper connector into their matching openings. Install retaining hardware lightly until alignment is confirmed.", (
    "Each control moves freely and remains accessible.",
    "Terminal polarity markings are visible and match the verified electrical documentation.",
    "No sharp edge bears against a component, wire, or insulation surface.",
))
step(doc, 8, "Fit the rear hardware", "Install the rear inlet/switch assembly and any ventilation or fan hardware only after its exact part and mounting method have been confirmed. Keep the computer PSU intact. Do not alter its internal mains circuitry.", (
    "The switch and inlet cannot rotate or pull through the panel.",
    "Required guards, barriers, and strain relief are present.",
    "Ventilation openings remain unobstructed.",
))
doc.add_page_break()
step(doc, 9, "Mount the internal modules", "Position the computer PSU, SK120, Arduino Nano, and relay on the intended mounting features. Begin with loose fasteners, check access and clearance, then tighten only enough to retain each part without cracking the PLA.", (
    "Fasteners cannot touch exposed circuitry.",
    "Modules do not contact one another or block cooling paths.",
    "Connectors remain accessible for the later verified wiring stage.",
))
step(doc, 10, "Plan cable routing", "Use the verified wiring schedule to plan paths before installing conductors. Keep wiring away from moving controls, sharp edges, hot parts, fan blades, and screw tips. Provide mechanical retention and preserve the required separation between hazardous and accessible circuits.", (
    "Every conductor endpoint will have a connector and pin reference.",
    "Wire color will not be the only means of identification.",
    "The assembly can be inspected before the enclosure is closed.",
))
step(doc, 11, "Complete the mechanical inspection", "Before any electrical work, check all mounting points, clearances, panel alignment, ventilation, and loose debris. Photograph the open assembly so later reviewers can see mounting and routing.", (
    "All modules and panels are securely retained.",
    "No loose screw, clipping, or conductive debris remains inside.",
    "The enclosure can close without pinching a cable.",
))

doc.add_heading("Mechanical build record", level=2)
record = doc.add_table(rows=1, cols=2)
record.alignment = WD_TABLE_ALIGNMENT.CENTER
record.autofit = False
for idx, text in enumerate(("Record", "Builder entry")):
    cell = record.rows[0].cells[idx]
    cell.width = Inches(2.1 if idx == 0 else 4.5)
    cell.text = text
    shade(cell, NAVY)
    margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
for row_number, label in enumerate(("Printer and nozzle", "PLA brand and color", "Slicer/profile", "Build date", "Hardware revision", "Fit adjustments or deviations"), start=1):
    cells = record.add_row().cells
    cells[0].text = label
    cells[1].text = ""
    for idx, cell in enumerate(cells):
        cell.width = Inches(2.1 if idx == 0 else 4.5)
        margins(cell, top=100, bottom=100)
        if row_number % 2 == 0:
            shade(cell, PALE_BLUE)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(8.8)
borders(record)

doc.add_page_break()
doc.add_heading("Electrical completion gate", level=1)
warning = doc.add_table(rows=1, cols=1)
cell = warning.cell(0, 0)
shade(cell, PALE_YELLOW)
margins(cell, top=140, bottom=140)
p = cell.paragraphs[0]
r = p.add_run("STOP BEFORE WIRING OR POWERING THE UNIT. ")
r.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)
p.add_run(
    "The current public release does not contain a verified conductor-by-conductor wiring schedule, exact module variants, final protective design, or commissioning limits. "
    "The preliminary schematic below records the intended architecture and marks unknown Nano and relay connections as TBC."
)

if SCHEMATIC_IMAGE.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(str(SCHEMATIC_IMAGE), width=Inches(6.85))
    p = doc.add_paragraph("Preliminary architecture only — dashed red connections must not be built until confirmed.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string("A61B1B")

doc.add_heading("Information required before electrical assembly", level=2)
for text in (
    "Exact computer PSU model, connector pinout, rail ratings, and condition assessment.",
    "Exact SK120, Arduino Nano, and relay variants with manufacturer documentation.",
    "Complete schematic and one wiring-schedule row for every conductor.",
    "Input protection, protective-earth/chassis arrangement, insulation, separation, strain relief, and enclosure suitability.",
    "Nano firmware, pin assignments, relay state during startup/reset/failure, and output-enable behavior.",
    "Defined inspection and test procedures with equipment, test points, numerical acceptance limits, and a responsible reviewer.",
):
    bullet(doc, text)

doc.add_heading("Final close-up and release checklist", level=1)
checks = (
    "All mechanical parts match one documented revision.",
    "Print settings and material use have been recorded.",
    "Every component and fastener has an exact specification and quantity.",
    "Internal assembly photographs show connector orientation, mounting, and cable routing.",
    "A qualified reviewer has approved the electrical design and protective measures.",
    "No-load, load, combined-load, display-accuracy, ripple, protection, and thermal results are published.",
    "A second person can reproduce the build using only the released files.",
)
for item in checks:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run("☐ ").bold = True
    p.add_run(item)

doc.core_properties.title = "Benchy Kira 03 Prototype Mechanical Assembly Guide"
doc.core_properties.subject = "Assembly guidance for the Benchy Kira 03 open-source enclosure prototype"
doc.core_properties.author = "Erdene Batbayar"
doc.core_properties.keywords = "Benchy, Kira 03, assembly guide, 3D printing, bench power supply"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
